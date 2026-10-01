"""
ocr.py - Document ingestion (Objective 2 / Scope 2.1 of the proposal).
Turns an uploaded PDF, DOCX, TXT or scanned image into plain text.

  PDF with a text layer  -> read directly (fast, exact)
  Scanned PDF / image    -> render pages to images and run Tesseract OCR (PyTesseract)
  DOCX                   -> read paragraphs and tables with python-docx
"""
import re
import shutil
from pathlib import Path

from . import config


class OCRUnavailable(Exception):
    """Raised when Tesseract is not installed, with a helpful message."""


def _setup_tesseract():
    import pytesseract
    if config.TESSERACT_CMD:
        pytesseract.pytesseract.tesseract_cmd = config.TESSERACT_CMD
    elif not shutil.which("tesseract"):
        win = Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe")
        if win.exists():
            pytesseract.pytesseract.tesseract_cmd = str(win)
    return pytesseract


def tesseract_status() -> str:
    """Used by the dashboard to show whether OCR is ready."""
    try:
        pt = _setup_tesseract()
        return "Ready (Tesseract " + str(pt.get_tesseract_version()) + ")"
    except Exception:
        return "Not installed - scanned documents cannot be read"


def _ocr_image(img):
    """OCR one PIL image -> (text, mean confidence 0-100)."""
    try:
        pt = _setup_tesseract()
        data = pt.image_to_data(img.convert("L"), output_type=pt.Output.DICT)
    except Exception as exc:
        name = exc.__class__.__name__
        if "TesseractNotFound" in name or "FileNotFound" in name:
            raise OCRUnavailable(
                "Tesseract OCR is not installed (or not found). Install it, or set TESSERACT_CMD in the .env file. "
                "See the guide, Part A, Step 7."
            )
        raise
    lines, confs = {}, []
    for i, word in enumerate(data["text"]):
        if not str(word).strip():
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        lines.setdefault(key, []).append(str(word))
        try:
            c = float(data["conf"][i])
            if c >= 0:
                confs.append(c)
        except (TypeError, ValueError):
            pass
    text = "\n".join(" ".join(v) for _, v in sorted(lines.items()))
    return text, (sum(confs) / len(confs) if confs else 0.0)


def clean_text(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(path: Path):
    """Return (text, method, ocr_confidence_or_None)."""
    ext = path.suffix.lower()

    if ext == ".txt":
        return clean_text(path.read_text(encoding="utf-8", errors="ignore")), "Plain text", None

    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs if p.text.strip()]
        for table in d.tables:
            for row in table.rows:
                parts.append("  ".join(c.text.strip() for c in row.cells))
        return clean_text("\n".join(parts)), "DOCX reader", None

    if ext == ".pdf":
        from pypdf import PdfReader
        try:
            reader = PdfReader(str(path))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[: config.MAX_OCR_PAGES])
        except Exception:
            text = ""
        if len(text.strip()) >= 60:
            return clean_text(text), "PDF text layer", None
        # No text layer -> it is a scan. Render each page to an image and OCR it.
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(str(path))
        texts, confs = [], []
        for i in range(min(len(pdf), config.MAX_OCR_PAGES)):
            img = pdf[i].render(scale=3).to_pil()
            t, c = _ocr_image(img)
            texts.append(t)
            confs.append(c)
        return clean_text("\n\n".join(texts)), "OCR (scanned PDF)", (sum(confs) / len(confs) if confs else 0.0)

    from PIL import Image
    img = Image.open(str(path))
    text, conf = _ocr_image(img)
    return clean_text(text), "OCR (image)", conf


def split_passages(text: str, words_per_chunk: int = None):
    """Cut text into short passages (used by semantic search)."""
    words_per_chunk = words_per_chunk or config.CHUNK_WORDS
    words = text.split()
    if not words:
        return []
    return [" ".join(words[i:i + words_per_chunk]) for i in range(0, len(words), words_per_chunk)]
